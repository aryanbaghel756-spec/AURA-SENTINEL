import os
import sys
import pptx
from pptx.util import Inches, Pt
import fitz # pymupdf
from PIL import Image

def assemble():
    template_path = r"C:\Users\aryan\OneDrive\Desktop\Aura\Docs\SIH2026-IDEA-Presentation-Format.pptx"
    if not os.path.exists(template_path):
        template_path = r"C:\Users\aryan\Downloads\SIH2026-IDEA-Presentation-Format.pptx"

    out_pptx = r"C:\Users\aryan\OneDrive\Desktop\AURA_SENTINEL_SIH2026_OFFICIAL_SUBMISSION.pptx"
    out_pdf = r"C:\Users\aryan\OneDrive\Desktop\AURA_SENTINEL_SIH2026_OFFICIAL_SUBMISSION.pdf"
    designer_dir = r"C:\Users\aryan\OneDrive\Desktop\Aura\Docs\designer_slides"

    prs = pptx.Presentation(template_path)
    print(f"Loaded template with {len(prs.slides)} slides.")

    # Width: 13.333 inches, Height: 7.5 inches (16:9)
    w_in = Inches(13.3333)
    h_in = Inches(7.5)

    # For each of the first 6 slides, clear child shapes and insert the full-bleed designer slide
    for idx in range(min(6, len(prs.slides))):
        slide = prs.slides[idx]
        slide_img = os.path.join(designer_dir, f"slide_{idx + 1}.png")
        if os.path.exists(slide_img):
            # Remove existing placeholder shapes
            for sp in list(slide.shapes):
                el = sp._element
                el.getparent().remove(el)
            # Add full-bleed high-res designer slide
            slide.shapes.add_picture(slide_img, 0, 0, w_in, h_in)
            print(f"Slide {idx + 1} assembled with high-res designer graphic.")

    # Delete slide 7 (Instructions slide) if present
    if len(prs.slides) >= 7:
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]
        print("Slide 7 deleted. Remaining count:", len(prs.slides))

    prs.save(out_pptx)
    print(f"Successfully saved PPTX to: {out_pptx}")

    # Generate perfect PDF directly from the 6 3000x1688 designer slides using PyMuPDF (fitz)
    # This guarantees 100% crispness, zero font-rendering glitches, zero COM locks!
    pdf_doc = fitz.open()
    for idx in range(6):
        slide_img = os.path.join(designer_dir, f"slide_{idx + 1}.png")
        img_doc = fitz.open(slide_img)
        rect = img_doc[0].rect
        pdfbytes = img_doc.convert_to_pdf()
        img_pdf = fitz.open("pdf", pdfbytes)
        page = pdf_doc.new_page(width=rect.width, height=rect.height)
        page.show_pdf_page(rect, img_pdf, 0)
        print(f"Added page {idx + 1} to PDF.")

    pdf_doc.save(out_pdf)
    pdf_doc.close()
    print(f"Successfully saved PDF to: {out_pdf}")

    # Verify PDF
    check_doc = fitz.open(out_pdf)
    print(f"VERIFIED FINAL PDF PAGE COUNT: {len(check_doc)}")
    check_doc.close()

if __name__ == '__main__':
    assemble()

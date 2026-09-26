import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_deck():
    template_path = r"C:\Users\aryan\Downloads\SIH2026-IDEA-Presentation-Format.pptx"
    out_pptx = r"C:\Users\aryan\OneDrive\Desktop\AURA_SENTINEL_SIH2026_OFFICIAL_SUBMISSION.pptx"
    assets_dir = r"C:\Users\aryan\OneDrive\Desktop\Aura\Docs\assets"
    qr_path = r"C:\Users\aryan\OneDrive\Desktop\Aura\Docs\github_qr.png"

    prs = pptx.Presentation(template_path)
    print(f"Loaded template with {len(prs.slides)} slides.")

    # Colors
    NAVY = RGBColor(15, 23, 42)
    DARK_BLUE = RGBColor(2, 132, 199)
    LIGHT_BG = RGBColor(248, 250, 252)
    BORDER_BLUE = RGBColor(186, 230, 253)
    GRAY_TEXT = RGBColor(100, 116, 139)
    WHITE = RGBColor(255, 255, 255)
    RED_ACCENT = RGBColor(225, 29, 72)
    GREEN_ACCENT = RGBColor(16, 185, 129)
    TEAL_ACCENT = RGBColor(15, 118, 110)
    PURPLE_ACCENT = RGBColor(124, 58, 237)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    s1 = prs.slides[0]
    for shape in list(s1.shapes):
        if shape.name in ["Picture 4", "Freeform: Shape 26"]:
            # remove placeholder graphic on right
            sp = shape._element
            sp.getparent().remove(sp)
        elif shape.name.startswith("Subtitle") and shape.has_text_frame:
            # Clear "TITLE PAGE" subtitle placeholder
            shape.text_frame.clear()
        elif shape.has_text_frame and "Problem Statement ID" in shape.text:
            tf = shape.text_frame
            tf.clear()

            fields = [
                ("Problem Statement ID:", " SIH26105"),
                ("Problem Statement Title:", " AI-Powered Continuous Cyber Risk Quantification & Investment Optimization Platform"),
                ("Theme:", " Blockchain & Cybersecurity"),
                ("PS Category:", " Software"),
                ("Team ID:", " 180219"),
                ("Team Name:", " ByteForce_1"),
                ("Team Leader:", " Aryan Baghel"),
                ("Team Members:", " Kirti, Hitesh Chauhan, Avinav Jha, Sharim Khan, Aman Sekh"),
                ("Institute:", " Skyline Institute of Engineering and Technology, Greater Noida")
            ]
            for idx, (label, val) in enumerate(fields):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.space_after = Pt(8)
                run1 = p.add_run()
                run1.text = label
                run1.font.bold = True
                run1.font.size = Pt(11.5)
                run1.font.color.rgb = NAVY
                run2 = p.add_run()
                run2.text = val
                run2.font.bold = False
                run2.font.size = Pt(11.5)
                run2.font.color.rgb = RGBColor(30, 41, 59)

    # Add Slide 1 Visual Graphic on the right
    s1_img = os.path.join(assets_dir, "slide1_visual.png")
    if os.path.exists(s1_img):
        s1.shapes.add_picture(s1_img, Inches(7.1), Inches(1.3), Inches(5.6), Inches(5.2))
    print("Slide 1 configured.")

    # -------------------------------------------------------------
    # Helper to clean default instruction boxes and set Team Oval
    # -------------------------------------------------------------
    def prepare_content_slide(slide, title_text):
        for shape in list(slide.shapes):
            if shape.name.startswith("Title") and shape.has_text_frame:
                shape.text_frame.text = title_text
                p = shape.text_frame.paragraphs[0]
                p.font.size = Pt(20)
                p.font.bold = True
                p.font.color.rgb = NAVY
            elif shape.name.startswith("Oval") and shape.has_text_frame:
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                r1 = p.add_run()
                r1.text = "ByteForce_1\n"
                r1.font.bold = True
                r1.font.size = Pt(9.5)
                r1.font.color.rgb = WHITE
                r2 = p.add_run()
                r2.text = "(ID: 180219)"
                r2.font.size = Pt(8.5)
                r2.font.color.rgb = RGBColor(224, 242, 254)
            elif shape.name.startswith("TextBox") and shape.has_text_frame:
                # Default instruction placeholder to remove
                sp = shape._element
                sp.getparent().remove(sp)

    # -------------------------------------------------------------
    # SLIDE 2: Idea Title & Proposed Solution
    # -------------------------------------------------------------
    s2 = prs.slides[1]
    prepare_content_slide(s2, "IDEA TITLE: AURA SENTINEL (AI-POWERED CYBER RISK QUANTIFICATION)")

    # Left Column: 3 Dark Rounded Cards (Lanezy Style)
    left_items = [
        ("Real-World Issue", 
         "Enterprises rely on static, qualitative 'High/Medium/Low' heatmaps that hide actual financial exposure, leaving boards unable to justify cyber budgets against dynamic threats.",
         RED_ACCENT, Inches(1.35)),
        ("Why Important",
         "Cyber incidents cost Indian enterprises ₹25,000+ Cr annually. 68% of CISOs struggle to defend security ROI to boards without quantified financial risk metrics.",
         DARK_BLUE, Inches(2.82)),
        ("Proposed Solution: AURA Sentinel",
         "AI-powered continuous cyber risk quantification combining the Open Group FAIR standard, Monte Carlo simulation, 0/1 Knapsack optimization, and SHA-256 Merkle blockchain.",
         GREEN_ACCENT, Inches(4.30))
    ]

    for title, desc, acc_col, y_pos in left_items:
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), y_pos, Inches(3.85), Inches(1.35))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(15, 23, 42)
        card.line.color.rgb = acc_col
        card.line.width = Pt(1.5)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.15)
        tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.1)
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.bold = True
        p0.font.size = Pt(10.5)
        p0.font.color.rgb = acc_col
        p0.space_after = Pt(3)
        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(8.5)
        p1.font.color.rgb = RGBColor(241, 245, 249)

    # Bottom Buttons on Left
    btn_x = Inches(0.55)
    for b_title in ["● Live SOC Console", "● GitHub Repo", "● REST API Docs"]:
        btn = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, btn_x, Inches(5.75), Inches(1.23), Inches(0.38))
        btn.fill.solid()
        btn.fill.fore_color.rgb = RGBColor(30, 41, 59)
        btn.line.color.rgb = DARK_BLUE
        btn.line.width = Pt(1)
        tf = btn.text_frame
        p = tf.paragraphs[0]
        p.text = b_title
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(7.5)
        p.font.bold = True
        p.font.color.rgb = RGBColor(56, 189, 248)
        btn_x += Inches(1.31)

    # Center Column: Pyramid Image
    s2_img = os.path.join(assets_dir, "slide2_pyramid.png")
    if os.path.exists(s2_img):
        s2.shapes.add_picture(s2_img, Inches(4.55), Inches(1.30), Inches(4.3), Inches(4.85))

    # Right Column: 3 Risk vs Solution Rows (Lanezy Style) with small "vs" badge
    rvs_items = [
        ("Static Annual Heatmaps", "Subjective High/Medium guesswork; stale within days; zero actuarial Rupee quantification.",
         "Continuous FAIR Engine", "Automated continuous loss calculation (Threat Frequency × Loss Magnitude in ₹).",
         Inches(1.35)),
        ("Arbitrary Budget Allocation", "Security budgets allocated via vendor hype & panic; no mathematical ROI justification.",
         "0/1 Knapsack Optimizer", "Bounded dynamic programming maximizes risk reduction per Rupee invested.",
         Inches(2.82)),
        ("Repudiable Audit Trails", "Editable database logs that fail rigorous regulatory scrutiny (DPDP & CERT-In).",
         "Merkle Blockchain Ledger", "Cryptographically sealed SHA-256 state anchoring for zero-trust compliance.",
         Inches(4.30))
    ]

    for r_title, r_desc, s_title, s_desc, y_pos in rvs_items:
        # Container
        cont = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.95), y_pos, Inches(3.85), Inches(1.35))
        cont.fill.solid()
        cont.fill.fore_color.rgb = LIGHT_BG
        cont.line.color.rgb = BORDER_BLUE
        cont.line.width = Pt(1)

        # Risk Box
        r_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.05), y_pos + Inches(0.06), Inches(3.65), Inches(0.52))
        r_box.fill.solid()
        r_box.fill.fore_color.rgb = RGBColor(254, 242, 242)
        r_box.line.color.rgb = RGBColor(254, 202, 202)
        tf_r = r_box.text_frame
        tf_r.word_wrap = True
        tf_r.margin_left = Inches(0.08)
        tf_r.margin_top = Inches(0.04)
        pr0 = tf_r.paragraphs[0]
        pr0.text = f"RISK: {r_title}"
        pr0.font.bold = True
        pr0.font.size = Pt(8.2)
        pr0.font.color.rgb = RED_ACCENT
        pr1 = tf_r.add_paragraph()
        pr1.text = r_desc
        pr1.font.size = Pt(7.0)
        pr1.font.color.rgb = NAVY

        # Small VS Badge in middle
        vs_badge = s2.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.65), y_pos + Inches(0.53), Inches(0.45), Inches(0.24))
        vs_badge.fill.solid()
        vs_badge.fill.fore_color.rgb = NAVY
        vs_badge.line.color.rgb = DARK_BLUE
        tf_vs = vs_badge.text_frame
        p_vs = tf_vs.paragraphs[0]
        p_vs.alignment = PP_ALIGN.CENTER
        p_vs.text = "vs"
        p_vs.font.size = Pt(6.5)
        p_vs.font.bold = True
        p_vs.font.color.rgb = WHITE

        # Solution Box
        s_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.05), y_pos + Inches(0.72), Inches(3.65), Inches(0.55))
        s_box.fill.solid()
        s_box.fill.fore_color.rgb = RGBColor(240, 253, 244)
        s_box.line.color.rgb = RGBColor(187, 247, 208)
        tf_s = s_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = Inches(0.08)
        tf_s.margin_top = Inches(0.04)
        ps0 = tf_s.paragraphs[0]
        ps0.text = f"SOLUTION: {s_title}"
        ps0.font.bold = True
        ps0.font.size = Pt(8.2)
        ps0.font.color.rgb = GREEN_ACCENT
        ps1 = tf_s.add_paragraph()
        ps1.text = s_desc
        ps1.font.size = Pt(7.0)
        ps1.font.color.rgb = NAVY

    print("Slide 2 configured.")

    # -------------------------------------------------------------
    # SLIDE 3: Technical Approach
    # -------------------------------------------------------------
    s3 = prs.slides[2]
    prepare_content_slide(s3, "TECHNICAL APPROACH & SYSTEM ARCHITECTURE")

    # Left: Circular Lifecycle Flowchart
    s3_left = os.path.join(assets_dir, "slide3_circular_flow.png")
    if os.path.exists(s3_left):
        s3.shapes.add_picture(s3_left, Inches(0.55), Inches(1.30), Inches(4.1), Inches(4.3))

    # Buttons below circular flow
    b1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(5.72), Inches(1.85), Inches(0.38))
    b1.fill.solid()
    b1.fill.fore_color.rgb = NAVY
    b1.line.color.rgb = DARK_BLUE
    tf1 = b1.text_frame
    p = tf1.paragraphs[0]
    p.text = "● Lifecycle Specification"
    p.font.size = Pt(8.0)
    p.font.bold = True
    p.font.color.rgb = WHITE

    b2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.65), Inches(5.72), Inches(1.85), Inches(0.38))
    b2.fill.solid()
    b2.fill.fore_color.rgb = NAVY
    b2.line.color.rgb = DARK_BLUE
    tf2 = b2.text_frame
    p = tf2.paragraphs[0]
    p.text = "● Algorithm Proofs"
    p.font.size = Pt(8.0)
    p.font.bold = True
    p.font.color.rgb = WHITE

    # Center: Architecture Diagram
    s3_center = os.path.join(assets_dir, "slide3_architecture.png")
    if os.path.exists(s3_center):
        s3.shapes.add_picture(s3_center, Inches(4.80), Inches(1.30), Inches(4.4), Inches(4.8))

    # Right: Technologies Used Stacked Cards
    tech_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.35), Inches(1.30), Inches(3.45), Inches(4.8))
    tech_box.fill.solid()
    tech_box.fill.fore_color.rgb = LIGHT_BG
    tech_box.line.color.rgb = BORDER_BLUE
    tf_tech = tech_box.text_frame
    tf_tech.word_wrap = True
    tf_tech.margin_left = Inches(0.14)
    tf_tech.margin_right = Inches(0.14)
    tf_tech.margin_top = Inches(0.12)

    p = tf_tech.paragraphs[0]
    p.text = "TECHNOLOGIES USED"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = NAVY
    p.space_after = Pt(6)

    tech_stacks = [
        ("Backend & APIs", "Python 3.10, FastAPI, Pydantic, Uvicorn, SQLite 3", DARK_BLUE),
        ("Quantification Engine", "Open Group FAIR (O-RT), NumPy, Monte Carlo (10k)", TEAL_ACCENT),
        ("Optimization Logic", "0/1 Knapsack (Dynamic Programming), Multi-Constraint", RGBColor(217, 119, 6)),
        ("Blockchain & Audit", "SHA-256 Merkle Tree, Cryptographic Hash Chain", PURPLE_ACCENT),
        ("Frontend & SOC UI", "React 18, Tailwind CSS, Lucide Icons, Recharts", RGBColor(37, 99, 235))
    ]

    for cat, items, col in tech_stacks:
        p_c = tf_tech.add_paragraph()
        p_c.text = f"▪ {cat}:"
        p_c.font.bold = True
        p_c.font.size = Pt(8.5)
        p_c.font.color.rgb = col
        p_c.space_after = Pt(1)

        p_i = tf_tech.add_paragraph()
        p_i.text = items
        p_i.font.size = Pt(8.0)
        p_i.font.color.rgb = NAVY
        p_i.space_after = Pt(4)

    print("Slide 3 configured.")

    # -------------------------------------------------------------
    # SLIDE 4: Feasibility and Viability
    # -------------------------------------------------------------
    s4 = prs.slides[3]
    prepare_content_slide(s4, "FEASIBILITY, VIABILITY & BUSINESS POTENTIAL")

    cards_s4 = [
        ("FEASIBILITY ANALYSIS",
         [("Plug-and-Play Ingestion", "Seamlessly connects with existing tools (Nessus, Qualys, AWS, Active Directory) via standard REST APIs without hardware changes."),
          ("Sub-Second Computation", "10,000 Monte Carlo runs execute in <850ms in pure NumPy; dynamic programming knapsack completes in <50ms."),
          ("Modular Microservices", "Dockerized, container-ready architecture allows independent scaling of analytics workers and API nodes."),
          ("Agentless Deployment", "Zero endpoint installation required, drastically reducing IT friction and operational overhead.")],
         DARK_BLUE, Inches(0.55)),

        ("VIABILITY & GOVERNANCE",
         [("Global Actuarial Standards", "Built strictly on Open Group FAIR & NIST SP 800-30 Rev 1, accepted by boards and cyber insurers globally."),
          ("Regulatory Compliance", "Merkle audit proofs satisfy DPDP Act 2023, RBI Cyber Security Framework, and CERT-In logging directives."),
          ("Boardroom Defensibility", "Converts ambiguous technical CVSS scores into CFO-ready ₹ Rupee loss distributions and justified ROI."),
          ("Zero Vendor Lock-In", "Open JSON/CSV data schemas and open-standards compliance ensure transparent, portable risk registers.")],
         TEAL_ACCENT, Inches(4.55)),

        ("BUSINESS POTENTIAL & ROI",
         [("70%+ Audit Cost Reduction", "Automates manual third-party cyber risk assessments, saving ₹20L–₹50L per annual enterprise audit cycle."),
          ("Optimized Security Spend", "Dynamic knapsack allocation prevents overspending on low-impact controls, saving 25–40% of security budget."),
          ("Enterprise Market Scale", "Targeted at BFSI, Healthcare, Critical Infrastructure, and IT enterprises needing continuous compliance."),
          ("Cyber Insurance Alignment", "Empirical Value-at-Risk modeling accelerates policy underwriting and lowers corporate insurance premiums.")],
         PURPLE_ACCENT, Inches(8.55))
    ]

    for c_title, points, col, x_pos in cards_s4:
        # Card container
        c_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos, Inches(1.30), Inches(3.85), Inches(4.85))
        c_box.fill.solid()
        c_box.fill.fore_color.rgb = LIGHT_BG
        c_box.line.color.rgb = col
        c_box.line.width = Pt(1.5)

        # Header pill
        h_pill = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_pos + Inches(0.1), Inches(1.40), Inches(3.65), Inches(0.42))
        h_pill.fill.solid()
        h_pill.fill.fore_color.rgb = col
        h_pill.line.color.rgb = col
        tf_h = h_pill.text_frame
        p = tf_h.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = c_title
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = WHITE

        # Content text box inside card
        tb = s4.shapes.add_textbox(x_pos + Inches(0.12), Inches(1.90), Inches(3.61), Inches(4.15))
        tf_c = tb.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = Inches(0.0)
        tf_c.margin_right = Inches(0.0)
        tf_c.margin_top = Inches(0.0)

        for p_idx, (p_title, p_desc) in enumerate(points):
            p_t = tf_c.paragraphs[0] if p_idx == 0 else tf_c.add_paragraph()
            p_t.text = f"✓ {p_title}"
            p_t.font.bold = True
            p_t.font.size = Pt(8.8)
            p_t.font.color.rgb = col
            p_t.space_after = Pt(1)

            p_d = tf_c.add_paragraph()
            p_d.text = p_desc
            p_d.font.size = Pt(7.8)
            p_d.font.color.rgb = NAVY
            p_d.space_after = Pt(5)

    print("Slide 4 configured.")

    # -------------------------------------------------------------
    # SLIDE 5: Impact and Benefits
    # -------------------------------------------------------------
    s5 = prs.slides[4]
    prepare_content_slide(s5, "IMPACT AND BENEFITS")

    # Center Hub / Left Area (Radiating Benefits Matrix)
    matrix_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.30), Inches(6.8), Inches(4.25))
    matrix_box.fill.solid()
    matrix_box.fill.fore_color.rgb = LIGHT_BG
    matrix_box.line.color.rgb = BORDER_BLUE
    matrix_box.line.width = Pt(1)

    # Hub header
    hub_header = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(1.42), Inches(6.4), Inches(0.42))
    hub_header.fill.solid()
    hub_header.fill.fore_color.rgb = NAVY
    hub_header.line.color.rgb = DARK_BLUE
    tf_hub = hub_header.text_frame
    p = tf_hub.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "AURA SENTINEL // MULTI-DIMENSIONAL IMPACT MATRIX"
    p.font.bold = True
    p.font.size = Pt(10)
    p.font.color.rgb = RGBColor(56, 189, 248)

    # Left operational benefits
    op_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(1.95), Inches(3.1), Inches(3.45))
    op_box.fill.solid()
    op_box.fill.fore_color.rgb = RGBColor(240, 253, 244)
    op_box.line.color.rgb = RGBColor(187, 247, 208)
    tf_op = op_box.text_frame
    tf_op.word_wrap = True
    tf_op.margin_left = Inches(0.1)
    tf_op.margin_right = Inches(0.1)
    tf_op.margin_top = Inches(0.08)

    p = tf_op.paragraphs[0]
    p.text = "OPERATIONAL ADVANTAGES"
    p.font.bold = True
    p.font.size = Pt(9.5)
    p.font.color.rgb = GREEN_ACCENT
    p.space_after = Pt(4)

    op_points = [
        ("24/7 Continuous Visibility", "Replaces periodic static audits with real-time dynamic risk quantification."),
        ("Optimal Budget ROI", "Knapsack algorithm maximizes security risk reduction per Rupee invested."),
        ("Impact Remediation", "Fixes vulnerabilities based on financial loss risk rather than CVSS noise alone."),
        ("Cryptographic Audit", "Tamper-proof Merkle ledger guarantees verifiable non-repudiation.")
    ]
    for opt, opd in op_points:
        p_t = tf_op.add_paragraph()
        p_t.text = f"● {opt}"
        p_t.font.bold = True
        p_t.font.size = Pt(8.2)
        p_t.font.color.rgb = RGBColor(6, 95, 70)
        p_d = tf_op.add_paragraph()
        p_d.text = opd
        p_d.font.size = Pt(7.2)
        p_d.font.color.rgb = NAVY
        p_d.space_after = Pt(3)

    # Right strategic pillars
    st_box = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.95), Inches(1.95), Inches(3.2), Inches(3.45))
    st_box.fill.solid()
    st_box.fill.fore_color.rgb = RGBColor(240, 249, 255)
    st_box.line.color.rgb = RGBColor(186, 230, 253)
    tf_st = st_box.text_frame
    tf_st.word_wrap = True
    tf_st.margin_left = Inches(0.1)
    tf_st.margin_right = Inches(0.1)
    tf_st.margin_top = Inches(0.08)

    p = tf_st.paragraphs[0]
    p.text = "STRATEGIC & REGULATORY PILLARS"
    p.font.bold = True
    p.font.size = Pt(9.5)
    p.font.color.rgb = DARK_BLUE
    p.space_after = Pt(4)

    st_points = [
        ("Boardroom Financial Clarity", "Translates cyber exposure into Value-at-Risk (₹ Crores) for executive consensus."),
        ("Compliance Readiness", "Automates reporting for DPDP Act 2023, RBI Cyber Guidelines, & CERT-In."),
        ("Cyber Insurance Underwriting", "Provides empirical loss exceedance curves required for favorable policy rates."),
        ("Unified C-Suite Vocabulary", "Bridges the technical communication gap between SOC analysts and CFO/Board.")
    ]
    for stt, std in st_points:
        p_t = tf_st.add_paragraph()
        p_t.text = f"● {stt}"
        p_t.font.bold = True
        p_t.font.size = Pt(8.2)
        p_t.font.color.rgb = RGBColor(3, 105, 161)
        p_d = tf_st.add_paragraph()
        p_d.text = std
        p_d.font.size = Pt(7.2)
        p_d.font.color.rgb = NAVY
        p_d.space_after = Pt(3)

    # Bottom Transformation Banner
    trans_banner = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(5.65), Inches(6.8), Inches(0.45))
    trans_banner.fill.solid()
    trans_banner.fill.fore_color.rgb = NAVY
    trans_banner.line.color.rgb = DARK_BLUE
    tf_tb = trans_banner.text_frame
    p = tf_tb.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "★ Paradigm Shift: From Subjective Guesswork to Actuarial Financial Precision ★"
    p.font.bold = True
    p.font.size = Pt(9.5)
    p.font.color.rgb = WHITE

    # Far Right: Performance Benchmark Chart
    s5_chart = os.path.join(assets_dir, "slide5_chart.png")
    if os.path.exists(s5_chart):
        s5.shapes.add_picture(s5_chart, Inches(7.55), Inches(1.30), Inches(5.2), Inches(4.8))

    print("Slide 5 configured.")

    # -------------------------------------------------------------
    # SLIDE 6: Research and References
    # -------------------------------------------------------------
    s6 = prs.slides[5]
    prepare_content_slide(s6, "RESEARCH, REFERENCES & LIVE DEMONSTRATION")

    # Top Section: 7 Reference Banner Cards (Lanezy Style)
    ref_items = [
        ("The Open Group FAIR Standard", "Risk Taxonomy (O-RT) & Risk Analysis (O-RA) for quantitative information risk modeling.", "opengroup.org"),
        ("NIST SP 800-30 Rev. 1", "Guide for Conducting Risk Assessments in Federal Information Systems.", "nist.gov"),
        ("ISO/IEC 27005:2022", "Information security, cybersecurity and privacy protection — Managing risks.", "iso.org"),
        ("DPDP Act 2023", "Digital Personal Data Protection Act, Ministry of Electronics & IT (MeitY), India.", "meity.gov.in"),
        ("CERT-In Directions (2022)", "Section 70B mandate for cybersecurity posture and log retention guidelines.", "cert-in.org.in"),
        ("0/1 Knapsack Optimization", "Dynamic Programming algorithms for resource-constrained risk minimization.", "ieee.org"),
        ("RFC 6962: Merkle Trees", "Verifiable append-only data structures for immutable security telemetry logging.", "rfc-editor.org")
    ]

    ref_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.30), Inches(11.8), Inches(2.25))
    ref_box.fill.solid()
    ref_box.fill.fore_color.rgb = LIGHT_BG
    ref_box.line.color.rgb = BORDER_BLUE
    ref_box.line.width = Pt(1)

    tf_ref = ref_box.text_frame
    tf_ref.word_wrap = True
    tf_ref.margin_left = Inches(0.12)
    tf_ref.margin_right = Inches(0.12)
    tf_ref.margin_top = Inches(0.08)

    p = tf_ref.paragraphs[0]
    p.text = "STANDARDS, FORMAL RESEARCH & REGULATORY CITATIONS"
    p.font.bold = True
    p.font.size = Pt(10)
    p.font.color.rgb = NAVY
    p.space_after = Pt(3)

    for name, desc, domain in ref_items:
        p_row = tf_ref.add_paragraph()
        p_row.space_after = Pt(2)
        r_name = p_row.add_run()
        r_name.text = f"• {name}: "
        r_name.font.bold = True
        r_name.font.size = Pt(8.0)
        r_name.font.color.rgb = DARK_BLUE

        r_desc = p_row.add_run()
        r_desc.text = f"{desc} "
        r_desc.font.size = Pt(7.5)
        r_desc.font.color.rgb = NAVY

        r_link = p_row.add_run()
        r_link.text = f"[{domain}]"
        r_link.font.bold = True
        r_link.font.size = Pt(7.5)
        r_link.font.color.rgb = GREEN_ACCENT

    # Bottom Left / Center: UI/UX Showcase Image
    s6_uiux = os.path.join(assets_dir, "slide6_uiux.png")
    if os.path.exists(s6_uiux):
        s6.shapes.add_picture(s6_uiux, Inches(0.55), Inches(3.68), Inches(8.3), Inches(2.45))

    # Bottom Right: GitHub QR Code Card
    qr_card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.0), Inches(3.68), Inches(3.35), Inches(2.45))
    qr_card.fill.solid()
    qr_card.fill.fore_color.rgb = NAVY
    qr_card.line.color.rgb = DARK_BLUE
    qr_card.line.width = Pt(1.5)

    tf_qr = qr_card.text_frame
    tf_qr.margin_left = Inches(0.1)
    tf_qr.margin_top = Inches(0.08)
    p = tf_qr.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "OFFICIAL CODE REPOSITORY"
    p.font.bold = True
    p.font.size = Pt(8.5)
    p.font.color.rgb = RGBColor(56, 189, 248)

    if os.path.exists(qr_path):
        s6.shapes.add_picture(qr_path, Inches(9.95), Inches(4.05), Inches(1.45), Inches(1.45))

    p_sub = tf_qr.add_paragraph()
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.space_before = Inches(1.55)
    r_sub = p_sub.add_run()
    r_sub.text = "ByteForce_1 // AURA SENTINEL\ngithub.com/aryanbaghel756-spec/AURA-SENTINEL"
    r_sub.font.size = Pt(7.0)
    r_sub.font.color.rgb = RGBColor(203, 213, 225)

    print("Slide 6 configured.")

    # -------------------------------------------------------------
    # SLIDE 7: Delete instruction slide (Ensure EXACTLY 6 slides)
    # -------------------------------------------------------------
    if len(prs.slides) >= 7:
        rId = prs.slides._sldIdLst[6].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[6]
        print("Slide 7 deleted. Remaining slides count:", len(prs.slides))

    prs.save(out_pptx)
    print(f"Successfully saved PPTX to: {out_pptx}")

if __name__ == '__main__':
    build_deck()
